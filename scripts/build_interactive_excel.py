import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def format_interactive_excel():
    # 1. Load your existing Excel workbook
    wb = openpyxl.load_workbook("Marketing_Analytics_Report.xlsx")

    # 2. Define Professional Styling Palette
    DARK_BLUE = "1B365D"
    ACCENT_BLUE = "2B547E"
    WHITE = "FFFFFF"
    GRAY_BORDER = "D9D9D9"
    ZEBRA_FILL = "F9FAFC"

    font_title = Font(name="Calibri", size=16, bold=True, color=WHITE)
    font_header = Font(name="Calibri", size=11, bold=True, color=WHITE)
    font_regular = Font(name="Calibri", size=11)

    fill_title = PatternFill(start_color=DARK_BLUE, end_color=DARK_BLUE, fill_type="solid")
    fill_header = PatternFill(start_color=ACCENT_BLUE, end_color=ACCENT_BLUE, fill_type="solid")
    fill_zebra = PatternFill(start_color=ZEBRA_FILL, end_color=ZEBRA_FILL, fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color=GRAY_BORDER),
        right=Side(style='thin', color=GRAY_BORDER),
        top=Side(style='thin', color=GRAY_BORDER),
        bottom=Side(style='thin', color=GRAY_BORDER)
    )

    # 3. Format 'ROI Summary' Sheet
    ws = wb["ROI Summary"]
    ws.insert_rows(1, 2)
    ws.merge_cells("A1:J1")
    ws["A1"] = "MARKETING CAMPAIGN PERFORMANCE & ROI DASHBOARD"
    ws["A1"].font = font_title
    ws["A1"].fill = fill_title
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    # Format Table Headers (Row 3)
    for col_idx in range(1, 11):
        cell = ws.cell(row=3, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Number Formats & Styling for Data Rows
    for row_idx in range(4, ws.max_row + 1):
        for col_idx in range(1, 11):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = font_regular
            cell.border = thin_border
            
            if row_idx % 2 == 0:
                cell.fill = fill_zebra
                
            # Number formatting
            if col_idx in [4, 5, 9]:  # Spend, Revenue, CAC
                cell.number_format = "₹#,##0.00"
                cell.alignment = Alignment(horizontal="right")
            elif col_idx in [6, 7]:   # Leads, Conversions
                cell.number_format = "#,##0"
                cell.alignment = Alignment(horizontal="right")
            elif col_idx == 8:        # Conv Rate
                cell.number_format = "0.00\"%\""
                cell.alignment = Alignment(horizontal="right")
            elif col_idx == 10:       # ROAS
                cell.number_format = "0.00\"x\""
                cell.alignment = Alignment(horizontal="right")

    # Auto-adjust Column Widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    # 4. Save formatted file
    wb.save("Interactive_Marketing_Analytics_Report.xlsx")
    print("Formatted report created: 'Interactive_Marketing_Analytics_Report.xlsx'")

if __name__ == "__main__":
    format_interactive_excel()